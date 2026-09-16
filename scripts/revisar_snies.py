import glob
import hashlib
import os
import re
from openpyxl import load_workbook

for ruta in sorted(glob.glob(os.path.join("data", "raw", "*.xlsx"))):
    nombre = os.path.basename(ruta)
    with open(ruta, "rb") as f:
        huella = hashlib.sha256(f.read()).hexdigest()
    try:
        wb = load_workbook(ruta, read_only=True)
        hojas = wb.sheetnames
        textos = []
        for fila in wb[hojas[0]].iter_rows(max_row=8, values_only=True):
            textos += [str(v).strip() for v in fila if v is not None and str(v).strip()]
        wb.close()
        resumen = " | ".join(textos)[:200]
        anios = sorted(set(re.findall(r"20[12]\d", resumen)))
        print(f"{nombre}  huella={huella[:16]}  hojas={hojas}  años={anios}\n    {resumen}\n")
    except Exception as e:
        print(f"{nombre}  huella={huella[:16]}  ERROR al abrir: {e}\n")